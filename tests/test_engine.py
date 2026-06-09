import os
import sys
import tempfile
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import engine


class EngineDuplicateSourceRowsTest(unittest.TestCase):
    def _make_target(self, feedback=197):
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        tmp.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = engine.TARGET_SHEET
        row = engine.TARGET_DATA_START_ROW
        ws.cell(row=row, column=engine.TGT["ID"] + 1).value = 219606014
        ws.cell(row=row, column=engine.TGT["2026反馈数量合计"] + 1).value = feedback
        wb.save(tmp.name)
        wb.close()
        return tmp.name

    def _source_row(self, inquiry=45, tm=2):
        return {
            engine.SRC["产品ID"]: 219606014,
            engine.SRC["询盘个数"]: inquiry,
            engine.SRC["TM咨询人数"]: tm,
        }

    def test_duplicate_source_ids_only_increment_feedback_once(self):
        target_path = self._make_target()
        try:
            source_rows = [self._source_row(), self._source_row(), self._source_row()]

            stats = engine.execute_update(target_path, source_rows)

            wb = openpyxl.load_workbook(target_path)
            ws = wb[engine.TARGET_SHEET]
            row = engine.TARGET_DATA_START_ROW
            feedback = ws.cell(row=row, column=engine.TGT["2026反馈数量合计"] + 1).value
            inquiry = ws.cell(row=row, column=engine.TGT["询盘数"] + 1).value
            tm = ws.cell(row=row, column=engine.TGT["TM数"] + 1).value
            wb.close()

            self.assertEqual(feedback, 244)
            self.assertEqual(inquiry, 45)
            self.assertEqual(tm, 2)
            self.assertEqual(stats.total, 3)
            self.assertEqual(stats.updated, 1)
            self.assertEqual(stats.deduplicated, 2)
            self.assertEqual(len(stats.warnings), 1)
        finally:
            os.unlink(target_path)

    def test_preview_matches_deduplicated_feedback_change(self):
        target_path = self._make_target()
        try:
            source_rows = [self._source_row(), self._source_row(), self._source_row()]

            stats = engine.analyze_changes(target_path, source_rows)

            feedback_changes = [
                change
                for detail in stats.changes
                for change in detail.changes
                if change[0] == "2026反馈数量合计"
            ]
            self.assertEqual(feedback_changes, [("2026反馈数量合计", 197, 244.0)])
            self.assertEqual(stats.deduplicated, 2)
        finally:
            os.unlink(target_path)


if __name__ == "__main__":
    unittest.main()
