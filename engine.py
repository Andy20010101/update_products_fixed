"""
Products 数据更新工具 - Core Engine
数据解析、更新、快照与撤销。
"""

import os
import re
import shutil
import glob as glob_mod
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

import xlrd
import openpyxl
from openpyxl.styles import PatternFill

RED_FILL = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")

# ============================================================
# 配置
# ============================================================
SOURCE_DIR = r"G:\阿里巴巴数据管家\阿里巴巴数据分析汇总\业务考核数据源和考核统计\产品数据表"
TARGET_DIR = r"\\192.168.1.10\zjh\社媒表汇总\社媒重要数据报表"
TARGET_FILENAME = "一：产品数据记录分析优化表_工作中.xlsx"
TARGET_SHEET = "良友效果好的产品变化"
SNAPSHOT_DIR_NAME = "snapshots"
SOURCE_HEADER_ROW = 6
TARGET_HEADER_ROW = 2
TARGET_DATA_START_ROW = 3

# Products 表列索引 (0-based)
SRC = {
    "产品ID": 0, "产品名称": 1, "是否橱窗": 2, "是否顶展": 3, "是否P4P": 4,
    "搜索曝光次数": 5, "搜索点击次数": 6, "搜索点击率": 7, "访问人数": 8,
    "询盘个数": 9, "询盘率": 12, "TM咨询人数": 17,
}

# 分析表列索引 (0-based)
TGT = {
    "ID": 3, "曝光数": 5, "点击数": 6, "点击率": 7, "访问人数": 9,
    "询盘数": 10, "询盘率": 11, "TM数": 13, "是否橱窗": 15, "是否顶展": 16,
    "是否P4P": 17, "2026反馈数量合计": 18,
}

PROTECTED_COLS = {0, 1, 2, 4, 8, 12, 14, 19, 20}

FIELD_MAP = [
    (SRC["搜索曝光次数"], TGT["曝光数"]),
    (SRC["搜索点击次数"], TGT["点击数"]),
    (SRC["搜索点击率"], TGT["点击率"]),
    (SRC["访问人数"], TGT["访问人数"]),
    (SRC["询盘个数"], TGT["询盘数"]),
    (SRC["询盘率"], TGT["询盘率"]),
    (SRC["TM咨询人数"], TGT["TM数"]),
    (SRC["是否橱窗"], TGT["是否橱窗"]),
    (SRC["是否顶展"], TGT["是否顶展"]),
    (SRC["是否P4P"], TGT["是否P4P"]),
]

RED_HIGHLIGHT_COLS = {TGT["曝光数"], TGT["点击数"], TGT["访问人数"], TGT["询盘数"], TGT["TM数"]}

NEW_PRODUCT_COLS = [
    TGT["ID"], TGT["曝光数"], TGT["点击数"], TGT["点击率"], TGT["访问人数"],
    TGT["询盘数"], TGT["询盘率"], TGT["TM数"], TGT["是否橱窗"], TGT["是否顶展"],
    TGT["是否P4P"], TGT["2026反馈数量合计"],
]

# 启动校验
_writable_cols = {tgt for _, tgt in FIELD_MAP} | {TGT["2026反馈数量合计"]} | set(NEW_PRODUCT_COLS)
_overlap = _writable_cols & PROTECTED_COLS
if _overlap:
    raise SystemExit(f"配置错误: 以下列同时在写入列和保护列中 (0-based): {_overlap}")


# ============================================================
# Data Models
# ============================================================

@dataclass
class SourceRow:
    """Products 源文件中的一行数据。"""
    product_id: str
    row_data: dict  # {列索引: 值}


@dataclass
class ChangeDetail:
    """单条变更详情。"""
    product_id: str
    action: str  # "更新" | "追加"
    changes: list = field(default_factory=list)  # [(字段名, 旧值, 新值), ...]


@dataclass
class UpdateStats:
    """更新统计。"""
    total: int = 0
    updated: int = 0
    appended: int = 0
    skipped: int = 0
    changes: list = field(default_factory=list)  # list[ChangeDetail]
    errors: list = field(default_factory=list)


# ============================================================
# Helper Functions
# ============================================================

def _convert_pct(val):
    if isinstance(val, str) and val.strip().endswith('%'):
        try:
            return float(val.strip().replace('%', '')) / 100.0
        except (ValueError, TypeError):
            return val
    return val


def _normalize_yn(val):
    if isinstance(val, str) and val.strip().upper() in ('Y', 'N'):
        return val.strip().lower()
    return val


VALUE_CONVERTERS = {
    TGT["点击率"]: _convert_pct,
    TGT["询盘率"]: _convert_pct,
    TGT["是否橱窗"]: _normalize_yn,
    TGT["是否顶展"]: _normalize_yn,
    TGT["是否P4P"]: _normalize_yn,
}


def normalize_id(val) -> str:
    try:
        f = float(str(val).strip())
        if f == int(f):
            return str(int(f))
        return str(f)
    except (ValueError, TypeError):
        return str(val).strip()


def safe_numeric(val, default=0):
    if val is None or val == "" or val == "-":
        return default
    try:
        v = float(str(val).replace("%", "").replace(",", "").strip())
        return v if v == v else default
    except (ValueError, TypeError):
        return default


def get_source_value(src_row: dict, src_col: int):
    val = src_row.get(src_col, "")
    if val is None:
        return ""
    return val


def _values_differ(old, new) -> bool:
    old_str = str(old).strip() if old is not None else ""
    new_str = str(new).strip() if new is not None else ""
    return old_str != new_str


# ============================================================
# File Discovery
# ============================================================

def find_latest_products_file(source_dir: str) -> Optional[str]:
    patterns = ["Products-*.xls", "Products-*.xlsx", "products-*.xls", "products-*.xlsx"]
    candidates = []
    for pat in patterns:
        candidates.extend(glob_mod.glob(os.path.join(source_dir, pat)))

    if not candidates:
        return None

    dated = []
    undated = []
    date_re = re.compile(r"(\d{4})-(\d{2})-(\d{2})")
    for f in candidates:
        m = date_re.search(os.path.basename(f))
        if m:
            dated.append((f, f"{m.group(1)}{m.group(2)}{m.group(3)}"))
        else:
            undated.append(f)

    if dated:
        dated.sort(key=lambda x: x[1], reverse=True)
        return dated[0][0]
    elif undated:
        undated.sort(key=lambda f: os.path.getmtime(f), reverse=True)
        return undated[0]
    return None


# ============================================================
# Source Parsing
# ============================================================

def read_products_source(filepath: str) -> list[dict]:
    """读取 Products 导出表，返回数据行列表。"""
    wb = xlrd.open_workbook(filepath)
    sheet = wb.sheet_by_index(0)

    header_row_idx = None
    config_idx = SOURCE_HEADER_ROW - 1
    if config_idx < sheet.nrows:
        cell_val = str(sheet.cell(config_idx, 0).value).strip()
        if "产品ID" in cell_val:
            header_row_idx = config_idx

    if header_row_idx is None:
        for rx in range(min(20, sheet.nrows)):
            cell_val = str(sheet.cell(rx, 0).value).strip()
            if "产品ID" in cell_val:
                header_row_idx = rx
                break

    if header_row_idx is None:
        raise ValueError("无法定位 Products 表标题行 (未找到「产品ID」列)")

    rows = []
    for rx in range(header_row_idx + 1, sheet.nrows):
        row = {}
        for cx in range(sheet.ncols):
            cell = sheet.cell(rx, cx)
            if cell.ctype == xlrd.XL_CELL_EMPTY:
                row[cx] = ""
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                if cell.value == int(cell.value):
                    row[cx] = int(cell.value)
                else:
                    row[cx] = cell.value
            else:
                row[cx] = str(cell.value).strip()
        if any(v != "" and v is not None for v in row.values()):
            rows.append(row)

    return rows


# ============================================================
# Analysis
# ============================================================

_SRC_COL_NAMES = {v: k for k, v in SRC.items()}
_TGT_COL_NAMES = {v: k for k, v in TGT.items()}


def _col_name(tgt_col: int) -> str:
    return _TGT_COL_NAMES.get(tgt_col, f"列{tgt_col}")


def _build_target_id_map(ws) -> dict:
    id_map = {}
    id_col = TGT["ID"] + 1
    for rx in range(TARGET_DATA_START_ROW, ws.max_row + 1):
        cell = ws.cell(row=rx, column=id_col)
        val = cell.value
        if val is not None and str(val).strip() != "":
            key = normalize_id(val)
            if key:
                id_map[key] = rx
    return id_map


def analyze_changes(target_path: str, source_rows: list[dict]) -> UpdateStats:
    """分析变更（只读），返回 UpdateStats 包含所有变更详情。"""
    wb = openpyxl.load_workbook(target_path)
    if TARGET_SHEET not in wb.sheetnames:
        raise ValueError(f"找不到 Sheet「{TARGET_SHEET}」，可用的 Sheet: {wb.sheetnames}")

    ws = wb[TARGET_SHEET]
    id_map = _build_target_id_map(ws)

    stats = UpdateStats(total=len(source_rows))
    id_col_src = SRC["产品ID"]

    for src_row in source_rows:
        src_id_raw = get_source_value(src_row, id_col_src)
        src_id = normalize_id(src_id_raw)
        if not src_id:
            stats.skipped += 1
            continue

        if src_id in id_map:
            target_row = id_map[src_id]
            changes = []
            for src_col, tgt_col in FIELD_MAP:
                val = get_source_value(src_row, src_col)
                converter = VALUE_CONVERTERS.get(tgt_col)
                if converter:
                    val = converter(val)
                if val != "" and val is not None:
                    cell = ws.cell(row=target_row, column=tgt_col + 1)
                    old_val = cell.value
                    if _values_differ(old_val, val):
                        changes.append((_col_name(tgt_col), old_val, val))

            # 2026反馈数量合计 变更
            inquiry = safe_numeric(get_source_value(src_row, SRC["询盘个数"]))
            tm = safe_numeric(get_source_value(src_row, SRC["TM咨询人数"]))
            existing = safe_numeric(ws.cell(row=target_row, column=TGT["2026反馈数量合计"] + 1).value)
            new_feedback = existing + inquiry + tm
            if new_feedback != existing:
                changes.append(("2026反馈数量合计", existing, new_feedback))

            if changes:
                stats.changes.append(ChangeDetail(product_id=src_id, action="更新", changes=changes))
            stats.updated += 1
        else:
            stats.changes.append(ChangeDetail(product_id=src_id, action="追加", changes=[]))
            stats.appended += 1

    wb.close()
    return stats


# ============================================================
# Update Execution
# ============================================================

def execute_update(target_path: str, source_rows: list[dict]) -> UpdateStats:
    """执行实际更新，保存到目标文件。返回统计信息。"""
    wb = openpyxl.load_workbook(target_path)
    if TARGET_SHEET not in wb.sheetnames:
        raise ValueError(f"找不到 Sheet「{TARGET_SHEET}」，可用的 Sheet: {wb.sheetnames}")

    ws = wb[TARGET_SHEET]
    id_map = _build_target_id_map(ws)

    stats = UpdateStats(total=len(source_rows))
    new_rows = []
    id_col_src = SRC["产品ID"]

    for src_row in source_rows:
        src_id_raw = get_source_value(src_row, id_col_src)
        src_id = normalize_id(src_id_raw)
        if not src_id:
            stats.skipped += 1
            continue

        if src_id in id_map:
            target_row = id_map[src_id]
            changes = []
            for src_col, tgt_col in FIELD_MAP:
                val = get_source_value(src_row, src_col)
                converter = VALUE_CONVERTERS.get(tgt_col)
                if converter:
                    val = converter(val)
                cell = ws.cell(row=target_row, column=tgt_col + 1)
                if val != "" and val is not None:
                    if tgt_col in RED_HIGHLIGHT_COLS and _values_differ(cell.value, val):
                        cell.fill = RED_FILL
                    if _values_differ(cell.value, val):
                        changes.append((_col_name(tgt_col), cell.value, val))
                    cell.value = val

            inquiry = safe_numeric(get_source_value(src_row, SRC["询盘个数"]))
            tm = safe_numeric(get_source_value(src_row, SRC["TM咨询人数"]))
            existing = safe_numeric(ws.cell(row=target_row, column=TGT["2026反馈数量合计"] + 1).value)
            feedback_cell = ws.cell(row=target_row, column=TGT["2026反馈数量合计"] + 1)
            new_feedback = existing + inquiry + tm
            if new_feedback != existing:
                changes.append(("2026反馈数量合计", existing, new_feedback))
            feedback_cell.value = new_feedback

            if changes:
                stats.changes.append(ChangeDetail(product_id=src_id, action="更新", changes=changes))
            stats.updated += 1
        else:
            new_row = {}
            new_row[TGT["ID"]] = get_source_value(src_row, SRC["产品ID"])
            new_row[TGT["曝光数"]] = get_source_value(src_row, SRC["搜索曝光次数"])
            new_row[TGT["点击数"]] = get_source_value(src_row, SRC["搜索点击次数"])
            new_row[TGT["点击率"]] = _convert_pct(get_source_value(src_row, SRC["搜索点击率"]))
            new_row[TGT["访问人数"]] = get_source_value(src_row, SRC["访问人数"])
            new_row[TGT["询盘数"]] = get_source_value(src_row, SRC["询盘个数"])
            new_row[TGT["询盘率"]] = _convert_pct(get_source_value(src_row, SRC["询盘率"]))
            new_row[TGT["TM数"]] = get_source_value(src_row, SRC["TM咨询人数"])
            new_row[TGT["是否橱窗"]] = _normalize_yn(get_source_value(src_row, SRC["是否橱窗"]))
            new_row[TGT["是否顶展"]] = _normalize_yn(get_source_value(src_row, SRC["是否顶展"]))
            new_row[TGT["是否P4P"]] = _normalize_yn(get_source_value(src_row, SRC["是否P4P"]))
            inquiry = safe_numeric(get_source_value(src_row, SRC["询盘个数"]))
            tm = safe_numeric(get_source_value(src_row, SRC["TM咨询人数"]))
            new_row[TGT["2026反馈数量合计"]] = inquiry + tm
            new_rows.append(new_row)
            stats.changes.append(ChangeDetail(product_id=src_id, action="追加", changes=[]))
            stats.appended += 1

    if new_rows:
        next_row = ws.max_row + 1
        for nr in new_rows:
            for tgt_col, val in nr.items():
                ws.cell(row=next_row, column=tgt_col + 1).value = val
            next_row += 1

    wb.save(target_path)
    return stats


# ============================================================
# Snapshot / Undo
# ============================================================

def create_snapshot(target_path: str, snapshot_dir: str) -> str:
    """创建目标文件的带时间戳快照，返回快照路径。"""
    os.makedirs(snapshot_dir, exist_ok=True)
    base, ext = os.path.splitext(os.path.basename(target_path))
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    snapshot_name = f"{base}_snapshot_{timestamp}{ext}"
    snapshot_path = os.path.join(snapshot_dir, snapshot_name)
    shutil.copy2(target_path, snapshot_path)
    return snapshot_path


def list_snapshots(snapshot_dir: str) -> list[str]:
    """列出所有快照，按文件名倒序（最新在前）。"""
    if not os.path.isdir(snapshot_dir):
        return []
    files = glob_mod.glob(os.path.join(snapshot_dir, "*_snapshot_*.xlsx"))
    files.sort(reverse=True)
    return files


def restore_from_snapshot(target_path: str, snapshot_path: str) -> None:
    """用指定快照覆盖目标文件。"""
    shutil.copy2(snapshot_path, target_path)


def cleanup_snapshots(snapshot_dir: str, keep: int = 20) -> int:
    """清理旧快照，只保留最近 keep 个。返回删除的数量。"""
    snapshots = list_snapshots(snapshot_dir)
    if len(snapshots) <= keep:
        return 0
    deleted = 0
    for sp in snapshots[keep:]:
        os.remove(sp)
        deleted += 1
    return deleted


def archive_source(source_path: str, source_dir: str) -> str:
    """将源文件移入 已处理 文件夹，返回归档路径。"""
    processed_dir = os.path.join(source_dir, "已处理")
    os.makedirs(processed_dir, exist_ok=True)
    dest = os.path.join(processed_dir, os.path.basename(source_path))
    if os.path.exists(dest):
        stem, ext = os.path.splitext(os.path.basename(source_path))
        dest = os.path.join(processed_dir, f"{stem}_{datetime.now():%Y%m%d_%H%M%S}{ext}")
    os.rename(source_path, dest)
    return dest
